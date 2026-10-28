import csv
import json
import math
import argparse
from pathlib import Path

from openpyxl import load_workbook

CARPETA = Path(__file__).resolve().parent
SALIDA = CARPETA.parents[1] / "public" / "data"

EXCEL_ENTRADA = CARPETA / "puntos-venta.xlsx"
CP_ENTRADA = CARPETA / "ES.txt"

PUNTOS_SALIDA = SALIDA / "puntos-venta.json"
CP_SALIDA = SALIDA / "codigos-postales.json"

NOMBRE_HOJA = None

DECIMALES = 5

def texto(valor):
    if valor is None:
        return ""

    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))

    return str(valor).strip()


def identificador(valor, longitud, permitir_vacio=False):
    resultado = texto(valor)

    if resultado == "":
        if permitir_vacio:
            return ""

        raise ValueError("Identificador obligatorio vacío.")

    if not resultado.isascii() or not resultado.isdigit():
        raise ValueError(f"Identificador no válido: {valor!r}")

    if len(resultado) > longitud:
        raise ValueError(
            f"Identificador demasiado largo: {resultado!r}; "
            f"máximo {longitud} dígitos."
        )

    return resultado.zfill(longitud)


def coordenada(valor, limite):
    cadena = texto(valor).replace(",", ".")

    if not cadena:
        raise ValueError("Coordenada vacía.")

    numero = float(cadena)

    if not math.isfinite(numero) or abs(numero) > limite:
        raise ValueError(f"Coordenada fuera de rango: {valor!r}")

    return numero

def convertir_excel():
    libro = load_workbook(
        EXCEL_ENTRADA,
        read_only=True,
        data_only=True
    )

    try:
        hoja = (
            libro[NOMBRE_HOJA]
            if NOMBRE_HOJA is not None
            else libro.worksheets[0]
        )

        filas = hoja.iter_rows(values_only=True)
        primera_fila = next(filas, None)

        if primera_fila is None:
            raise ValueError("El Excel está vacío.")

        cabeceras = [
            texto(valor).lstrip("\ufeff").upper()
            for valor in primera_fila
        ]

        requeridas = [
            "DESPACHO",
            "ADMON",
            "DOMICILIO",
            "CODPOSTAL",
            "LOCALIDAD",
            "LATITUD",
            "LONGITUD",
        ]

        faltantes = [
            nombre for nombre in requeridas
            if nombre not in cabeceras
        ]

        if faltantes:
            raise ValueError(
                "Faltan columnas en el Excel: "
                + ", ".join(faltantes)
            )

        indices = {
            nombre: cabeceras.index(nombre)
            for nombre in requeridas
        }

        puntos = []

        for numero_fila, fila in enumerate(filas, start=2):
            if all(texto(valor) == "" for valor in fila):
                continue

            def obtener(nombre):
                return fila[indices[nombre]]

            try:
                puntos.append([
                    identificador(obtener("DESPACHO"), 5),
                    identificador(
                        obtener("ADMON"),
                        9,
                        permitir_vacio=True
                    ),
                    texto(obtener("DOMICILIO")),
                    identificador(obtener("CODPOSTAL"), 5),
                    texto(obtener("LOCALIDAD")),
                    round(coordenada(obtener("LATITUD"), 90), DECIMALES),
                    round(coordenada(obtener("LONGITUD"), 180), DECIMALES),
                ])
            except (ValueError, TypeError) as error:
                raise ValueError(
                    f"Excel, fila {numero_fila}: {error}"
                ) from error

        guardar_json(PUNTOS_SALIDA, puntos)

        return len(puntos)

    finally:
        libro.close()

def convertir_codigos_postales():
    grupos = {}
    filas_leidas = 0

    with CP_ENTRADA.open(
        encoding="utf-8-sig", newline=""
    ) as entrada:

        lector = csv.reader(entrada, delimiter="\t")

        for numero_fila, campos in enumerate(lector, start=1):
            if not campos or all(
                not campo.strip() for campo in campos
            ):
                continue

            if len(campos) != 12:
                raise ValueError(
                    f"ES.txt, fila {numero_fila}: "
                    "se esperaban 12 columnas separadas "
                    f"por tabuladores y se encontraron {len(campos)}."
                )

            if campos[0].strip() != "ES":
                raise ValueError(
                    f"ES.txt, fila {numero_fila}: "
                    "el código de país no es ES."
                )

            try:
                cp = identificador(campos[1], 5)
                latitud = coordenada(campos[9], 90)
                longitud = coordenada(campos[10], 180)
            except (ValueError, TypeError) as error:
                raise ValueError(
                    f"ES.txt, fila {numero_fila}: {error}"
                ) from error

            grupos.setdefault(cp, set()).add((latitud, longitud))
            filas_leidas += 1

    def distancia_km(punto_a, punto_b):
        lat_a, lng_a = punto_a
        lat_b, lng_b = punto_b

        lat_a = math.radians(lat_a)
        lat_b = math.radians(lat_b)

        delta_lat = lat_b - lat_a
        delta_lng = math.radians(lng_b - lng_a)

        h = (
            math.sin(delta_lat / 2) ** 2
            + math.cos(lat_a)
            * math.cos(lat_b)
            * math.sin(delta_lng / 2) ** 2
        )

        h = max(0.0, min(1.0, h))

        return 6371 * 2 * math.asin(math.sqrt(h))

    codigos = []
    cp_con_varios_puntos = 0

    for cp, coordenadas_unicas in sorted(grupos.items()):
        puntos = sorted(coordenadas_unicas)

        if len(puntos) == 1:
            latitud, longitud = puntos[0]
        else:
            cp_con_varios_puntos += 1

            centro = (
                math.fsum(lat for lat, lng in puntos) / len(puntos),
                math.fsum(lng for lat, lng in puntos) / len(puntos)
            )

            latitud, longitud = min(
                puntos,
                key=lambda punto: distancia_km(punto, centro)
            )

        codigos.append([
            cp,
            round(longitud, DECIMALES),
            round(latitud, DECIMALES)
        ])

    guardar_json(CP_SALIDA, codigos)

    print(f"Filas de origen: {filas_leidas}")
    print(f"Códigos postales únicos: {len(codigos)}")
    print(
        "CP con varias coordenadas distintas: "
        f"{cp_con_varios_puntos}"
    )

    return len(codigos)

def guardar_json(ruta, datos):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_text(
        json.dumps(
            datos,
            ensure_ascii=False,
            separators=(",", ":"),
            allow_nan=False
        ),
        encoding="utf-8"
    )

def mostrar_resumen(ruta, registros):
    print(
        f"{ruta.name}: {registros} registros, "
        f"{ruta.stat().st_size / 1024:.1f} KB"
    )

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Genera los JSON compactos del buscador."
    )

    parser.add_argument(
        "tipo",
        choices=["pv", "cp"],
        help=(
            "pv: genera puntos-venta.json desde el Excel; "
            "cp: genera codigos-postales.json desde ES.txt."
        )
    )

    args = parser.parse_args()

    if args.tipo == "pv":
        registros = convertir_excel()
        mostrar_resumen(PUNTOS_SALIDA, registros)

    elif args.tipo == "cp":
        registros = convertir_codigos_postales()
        mostrar_resumen(CP_SALIDA, registros)